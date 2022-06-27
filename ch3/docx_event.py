import openpyxl as excel, docx, os

# 설정
in_file = './input/name_addr.xlsx'
template_file = './input/event-template.docx'
save_dir = os.path.join(os.path.dirname(__file__),'output', 'events')

# 안내장을 저장할 폴더 생성 --- (1)
if not os.path.exists(save_dir):
    os.mkdir(save_dir)

# 엑셀 파일을 열고 행 데이터를 리스트로 가져오기 --- (2)
def read_book():
    result = []
    sheet = excel.load_workbook(in_file).active
    for row in sheet.iter_rows(min_row=2):
        v = [c.value for c in row]
        if v[0] is None: break
        result.append(v)
    return result

# 고객의 수만큼 안내장 생성 --- (3)
for person in read_book():
    name, zipnum, addr = person
    card = {
        '[[주소]]': '(우)'+zipnum+' | '+addr,
        '[[고객명]]': name
    }
    # 워드 템플릿 읽기 --- (4)
    doc = docx.Document(template_file)
    # 내용을 바꿔 쓰기
    for p in doc.paragraphs:
        # 텍스트 교체 --- (5)
        for k,v in card.items():
            if p.text.find(k) >= 0:
                p.text = p.text.replace(k, v)
                p.runs[0].font.bold = True
    # 워드 파일 저장 --- (6)
    save_file = os.path.join(save_dir,
                  name+' 님.docx')
    print('save:', save_file)
    doc.save(save_file)
